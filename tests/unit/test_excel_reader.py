"""Tests for Excel ingestion."""

import pytest

from business_analyst.exceptions import (
    DuplicateHeaderError,
    EmptyDatasetError,
    InvalidFileError,
    RowLimitError,
    SheetSelectionError,
)
from business_analyst.ingestion.excel_reader import (
    list_excel_sheets,
    read_excel_bytes,
)


def test_lists_excel_sheets(xlsx_bytes) -> None:
    assert list_excel_sheets(xlsx_bytes) == ["Sales", "Preview"]


def test_reads_selected_worksheet(xlsx_bytes, test_settings) -> None:
    dataframe = read_excel_bytes(
        xlsx_bytes,
        "Sales",
        test_settings,
    )

    assert dataframe.shape == (2, 2)
    assert list(dataframe.columns) == ["order_id", "revenue"]


def test_rejects_missing_worksheet(xlsx_bytes, test_settings) -> None:
    with pytest.raises(SheetSelectionError):
        read_excel_bytes(xlsx_bytes, "Missing", test_settings)


def test_rejects_invalid_workbook(test_settings) -> None:
    with pytest.raises(InvalidFileError):
        list_excel_sheets(b"not-an-excel-file")


def test_rejects_header_only_worksheet(test_settings) -> None:
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["name", "revenue"])

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    with pytest.raises(EmptyDatasetError):
        read_excel_bytes(
            buffer.getvalue(),
            "Sheet",
            test_settings,
        )


def test_rejects_duplicate_excel_headers(test_settings) -> None:
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["revenue", "revenue"])
    worksheet.append([100, 200])

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    with pytest.raises(DuplicateHeaderError):
        read_excel_bytes(
            buffer.getvalue(),
            "Sheet",
            test_settings,
        )


def test_rejects_excess_excel_rows(test_settings) -> None:
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["name", "value"])

    for index in range(6):
        worksheet.append([f"Row-{index}", index])

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    with pytest.raises(RowLimitError):
        read_excel_bytes(
            buffer.getvalue(),
            "Sheet",
            test_settings,
        )
