"""Excel workbook inspection and worksheet parsing."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from business_analyst.config import Settings
from business_analyst.exceptions import (
    ColumnLimitError,
    DuplicateHeaderError,
    EmptyDatasetError,
    InvalidFileError,
    MissingHeaderError,
    RowLimitError,
    SheetSelectionError,
)
from business_analyst.ingestion.file_checks import validate_xlsx_signature


def list_excel_sheets(data: bytes) -> list[str]:
    """Return worksheet names after basic workbook validation."""
    validate_xlsx_signature(data)

    try:
        workbook = load_workbook(
            BytesIO(data),
            read_only=True,
            data_only=True,
        )
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()
    except Exception as exc:
        raise InvalidFileError("The Excel workbook could not be opened.") from exc


def _validate_headers(
    columns: list[object],
    max_columns: int,
) -> list[str]:
    """Validate worksheet column labels."""
    if not columns or all(
        pd.isna(column) or not str(column).strip() for column in columns
    ):
        raise MissingHeaderError("The selected worksheet has no usable column headers.")

    if len(columns) > max_columns:
        raise ColumnLimitError(
            f"The worksheet exceeds the limit of {max_columns} columns."
        )

    normalized = [str(column).strip() for column in columns]

    if any(not column for column in normalized):
        raise MissingHeaderError("One or more worksheet headers are empty.")

    if len(normalized) != len(set(normalized)):
        raise DuplicateHeaderError(
            "The selected worksheet contains duplicate column headers."
        )

    return normalized


def _read_original_headers(
    data: bytes,
    sheet_name: str,
) -> list[object]:
    """Read original worksheet header cells using openpyxl."""
    workbook = load_workbook(
        BytesIO(data),
        read_only=True,
        data_only=True,
    )

    try:
        worksheet = workbook[sheet_name]
        first_row = next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            ),
            None,
        )
        return list(first_row) if first_row is not None else []
    finally:
        workbook.close()


def read_excel_bytes(
    data: bytes,
    sheet_name: str,
    settings: Settings,
) -> pd.DataFrame:
    """Read a selected worksheet with configured row and column limits."""
    sheets = list_excel_sheets(data)

    if not sheets:
        raise InvalidFileError("The Excel workbook does not contain any worksheets.")

    if sheet_name not in sheets:
        raise SheetSelectionError(
            "The requested worksheet does not exist in this workbook."
        )

    # Validate the original header before pandas can mangle duplicate names.
    # Keep these checks outside the parser exception wrapper so the original
    # application exception types are preserved.
    original_headers = _read_original_headers(data, sheet_name)
    columns = _validate_headers(original_headers, settings.max_columns)

    try:
        dataframe = pd.read_excel(
            BytesIO(data),
            sheet_name=sheet_name,
            engine="openpyxl",
            nrows=settings.max_rows + 1,
        )
    except (pd.errors.ParserError, UnicodeError, ValueError, KeyError) as exc:
        raise InvalidFileError("The selected worksheet could not be parsed.") from exc
    except Exception as exc:
        raise InvalidFileError("The selected worksheet could not be read.") from exc

    if len(dataframe) > settings.max_rows:
        raise RowLimitError(
            f"The worksheet exceeds the limit of {settings.max_rows} data rows."
        )

    if dataframe.empty:
        raise EmptyDatasetError(
            "The selected worksheet contains headers but no data rows."
        )

    if len(dataframe.columns) != len(columns):
        raise InvalidFileError(
            "The worksheet header changed unexpectedly during parsing."
        )

    # Use the validated, normalized original labels as the DataFrame columns.
    dataframe.columns = columns
    return dataframe
